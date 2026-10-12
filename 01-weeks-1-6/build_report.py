import json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parent
WALLET = json.loads((ROOT / "evidence" / "wallet.json").read_text())
DETECTIVE_1 = json.loads((ROOT / "evidence" / "detective" / "transaction-1.json").read_text())
DETECTIVE_2 = json.loads((ROOT / "evidence" / "detective" / "transaction-2.json").read_text())
CLIENT = json.loads((ROOT / "evidence" / "w4-w6" / "client_verification.json").read_text())

SECTIONS = [
    (
        "Сводка выполнения",
        [
            "Работа подготовлена по курсу ИТ-стартап Итоговый проект 1 для группы Data Science 3 курса. Преподаватель: Иманкулова Алия Жангельдиновна. В проект собраны материалы по шести обязательным практическим заданиям: вход в Web3, разбор транзакций Transaction Detective, CLI-программа на Rust, первая программа Solana, счётчик на Anchor с клиентской проверкой и профиль пользователя на PDA.",
            "Важное исправление к прежней отправке от 1 октября: в том файле за перевод студента была ошибочно выдана чужая транзакция Devnet, а вывод локального симулятора был описан как работа в сети. В настоящем отчёте все результаты разделены честно: где запросы шли в живой Devnet RPC, где разбирались публичные транзакции других участников и где выполнялись локальные тесты программ.",
        ],
        None,
    ),
    (
        "Задание 1. Кошелёк, тестовый SOL и первая транзакция",
        [
            f"Для работы созданы два настоящих Ed25519-кошелька Solana: основной адрес студента {WALLET['student']} и адрес получателя {WALLET['recipient']}. Приватные ключи сохранены только локально с правами 600 и не включены ни в отчёт, ни в архив.",
            f"Через публичный узел {WALLET['rpc']} проверен баланс на слоте {WALLET['balance']['context']['slot']}: на счёте 0 lamports. Автоматические запросы requestAirdrop к публичному RPC и через solfaucet.com вернули ошибки Internal error и код 429 о лимите крана, а веб-краны faucet.solana.com и devnetfaucet.org требуют ручного прохождения Cloudflare или входа через GitHub в браузере студента.",
            "Из-за нулевого баланса перевод в публичном Devnet пока не отправлен. Скрипт отправки devnet_work.py полностью готов: как только на адрес студента поступит тестовый SOL, он подпишет и отправит реальный SystemProgram::Transfer. Дополнительно тот же перевод проверен в локальной виртуальной машине LiteSVM: после пополнения тестового счёта на 2 SOL перевод 100000000 lamports завершился со статусом success, баланс получателя стал равен 100000000 lamports.",
        ],
        None,
    ),
    (
        "Задание 2. Transaction Detective: разбор реальных транзакций",
        [
            "Для второго задания через метод getTransaction с параметром jsonParsed получены две настоящие подтверждённые транзакции публичной сети Devnet. Они принадлежат сторонним участникам сети и используются только как объект анализа.",
            f"Первая транзакция: {DETECTIVE_1['signature']} в слоте {DETECTIVE_1['result']['slot']}. Плательщик комиссии: {DETECTIVE_1['result']['transaction']['message']['accountKeys'][0]['pubkey']}, списано {DETECTIVE_1['result']['meta']['fee']} lamports, ошибка выполнения отсутствует. Внутри пакета 6 инструкций: две инструкции ComputeBudget111111111111111111111111111111 задают лимит вычислений и приоритетную комиссию, программа MemoSq4gqABAXKb96qnH8TysNcWxMyWCqXgDLGmfcHr записывает строку метаданных сделки, два вызова ATokenGPvbdGVxr1b2hvZbsiqW5xWH25efTNsLJA8knL выполняют createIdempotent для токен-аккаунтов, после чего TokenkegQfeZyiNwAJbNbGKPFXCWuBvf9Ss623VQ5DA переводит 1000000000 базовых единиц токена mint 4wMp7FQJrSwHXJ92kKRd9X2LQE1Tw7L4beT7fBLgjY4y.",
            f"Вторая транзакция: {DETECTIVE_2['signature']} в слоте {DETECTIVE_2['result']['slot']}. Плательщик комиссии: {DETECTIVE_2['result']['transaction']['message']['accountKeys'][0]['pubkey']}, комиссия {DETECTIVE_2['result']['meta']['fee']} lamports. После двух инструкций Compute Budget и записи в SPL Memo создаётся связанный токен-аккаунт DFJbHB9qhxWm7QeZ5DS478j6YHBZHbzJaTG7wk5FRCCh для mint Aj1qyqJRvi7VFdCriNvLiiCvjkLW3pTRgFCKSyGtAbsk, выполняется перевод 69511888 единиц токена и вызывается пользовательская программа Bz4AoAXL8vkTSFPqR52GwNxFixAbKPJ3VShhgoS2eV2G. Поскольку её IDL не опубликован в RPC-ответе, данные инструкции оставлены в исходном Base58 виде без домыслов о названии метода.",
            "Разбор подтверждает ключевой принцип архитектуры Solana: исполняемый код программ отделён от аккаунтов данных, права подписи и изменения задаются флагами signer и writable в заголовке сообщения, а перевод SPL-токена меняет поля токен-аккаунтов, но не является системным переводом SOL.",
        ],
        None,
    ),
    (
        "Задание 3. CLI-программа на Rust для разработчика Solana",
        [
            "В каталоге rust-cli написана и собрана консольная утилита solana-lab-cli на Rust. В отличие от старого макета она не печатает зашитые в код хеши, а проверяет длину адресов и сигнатур по алфавиту Base58 и обращается к настоящему JSON-RPC узлу https://api.devnet.solana.com.",
            "В программе использованы структуры, перечисления команд и ошибок, сопоставление с образцом match, обработка ошибок через Result и модульные тесты. Команда cargo test проходит все 5 проверок: адрес System Program декодируется ровно в 32 байта, строка с запрещённым символом l отклоняется, короткая сигнатура отклоняется, запрос баланса для валидного ключа проходит.",
        ],
        "cargo test --manifest-path rust-cli/Cargo.toml\nrunning 5 tests\ntest tests::invalid_chars ... ok\ntest tests::short_signature ... ok\ntest tests::system_address ... ok\ntest tests::usage ... ok\ntest tests::valid_balance ... ok\n\ncargo run --manifest-path rust-cli/Cargo.toml -- balance 2SzJ2PSuhnQsenfvHK4DhQHibjVL2d68fwdHAMwxJ7pL\n{\"jsonrpc\":\"2.0\",\"result\":{\"context\":{\"apiVersion\":\"4.4.0-beta.0\",\"slot\":510169500},\"value\":0},\"id\":1}",
    ),
    (
        "Задание 4. Первая программа Solana",
        [
            "В каталоге programs/first-program реализована нативная программа на крейте solana-program 2.3.0 с точкой входа entrypoint!(process_instruction). Программа принимает байтовый буфер инструкции: при передаче байта 0 она записывает приветствие в журнал выполнения через макрос msg! и возвращает Ok(()), а для любых других входных данных возвращает ошибку InvalidInstructionData.",
            "Модульные тесты проверяют успешный вызов и отказ на неизвестной инструкции. Для реального деплоя в Devnet требуются установленный компилятор cargo-build-sbf и пополненный тестовый кошелёк; на текущем сервере свободно около 108 МБ дискового пространства, поэтому тяжёлый SBF SDK сюда не помещается, а кран Devnet закрыт лимитом. Сама программа и тесты полностью готовы к сборке и выкладке после пополнения кошелька.",
        ],
        "pub fn process_instruction(_program_id: &Pubkey, _accounts: &[AccountInfo], data: &[u8]) -> ProgramResult {\n    match data {\n        [0] => { msg!(\"Hello from Emir's first Solana program!\"); Ok(()) },\n        _ => Err(solana_program::program_error::ProgramError::InvalidInstructionData),\n    }\n}",
    ),
    (
        "Задание 5. Программа счётчика на Anchor и клиентская проверка",
        [
            f"В каталоге programs/anchor-counter написан контракт anchor_counter на базе anchor-lang с объявленным идентификатором {CLIENT['counter_program_id']}. Аккаунт состояния Counter хранит публичный ключ владельца authority и беззнаковый счётчик count. Реализованы три инструкции: initialize создаёт аккаунт со значением 0, increment увеличивает счётчик с защитой checked_add, decrement уменьшает значение с проверкой checked_sub.",
            f"Ограничение #[account(mut, has_one = authority)] гарантирует, что изменить счётчик может только владелец, подписавший транзакцию. В модульном тесте через точку входа Anchor entry выполнены два инкремента до значения 2, проверен отказ при попытке вызова от чужого подписанта, затем выполнены два декремента до 0 и проверен отказ при попытке уйти в отрицательное значение. Клиентский скрипт clients/solana_client.py вычисляет 8-байтовый дискриминатор инструкции increment ({CLIENT['counter_increment_discriminator_hex']}) и проверяет бинарную сериализацию аккаунта Counter по схеме Borsh.",
        ],
        "running 2 tests\ntest test_id ... ok\ntest tests::counter_client_roundtrip_and_authority_check ... ok",
    ),
    (
        "Задание 6. Децентрализованный профиль на PDA",
        [
            f"В каталоге programs/pda-profile реализована программа pda_profile с идентификатором {CLIENT['profile_program_id']}. Профиль пользователя размещается по детерминированному адресу Program Derived Address со связкой семян [b\"profile\", authority.key().as_ref()] и каноническим bump. Для ключа студента {CLIENT['student_authority']} клиентом и программой выводится один и тот же адрес PDA {CLIENT['derived_profile_pda']} с bump={CLIENT['derived_profile_bump']}.",
            "Инструкция create_profile инициализирует аккаунт UserProfile и сохраняет поля authority, username (до 32 байт), bio (до 160 байт) и bump. Инструкция update_profile проверяет подпись владельца через has_one = authority и совпадение PDA по сохранённому bump. В тесте программы проверены детерминированный вывод адреса, обновление профиля на значения emir_ds и AlmaU 3 course, отказ при подписи чужим ключом и отказ при превышении длины имени пользователя.",
        ],
        "running 2 tests\ntest test_id ... ok\ntest tests::pda_derivation_update_and_owner_check ... ok",
    ),
]

def build_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(10.5)
    head = doc.add_paragraph()
    r = head.add_run("Яппаров Эмир DS")
    r.bold = True
    r.font.size = Pt(12)
    title = doc.add_paragraph()
    tr = title.add_run("ИТ-стартап (Solana): отчёт по практическим заданиям 1-6")
    tr.bold = True
    tr.font.size = Pt(14)
    for heading, paragraphs, code in SECTIONS:
        h = doc.add_paragraph()
        hr = h.add_run(heading)
        hr.bold = True
        hr.font.size = Pt(12)
        for text in paragraphs:
            p = doc.add_paragraph(text)
            p.paragraph_format.space_after = Pt(6)
        if code:
            cp = doc.add_paragraph()
            cr = cp.add_run(code)
            cr.font.name = "DejaVu Sans Mono"
            cr.font.size = Pt(9)
    doc.save(path)

def build_pdf(path: Path) -> None:
    pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("DejaVuBold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
    pdfmetrics.registerFont(TTFont("DejaVuMono", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["Normal"], fontName="DejaVu", fontSize=10, leading=14, spaceAfter=6)
    h1 = ParagraphStyle("H1", parent=body, fontName="DejaVuBold", fontSize=13, leading=17, spaceBefore=8, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=body, fontName="DejaVuBold", fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=4)
    code_style = ParagraphStyle("Code", fontName="DejaVuMono", fontSize=8.3, leading=11, leftIndent=8, spaceAfter=6)
    story = [
        Paragraph("Яппаров Эмир DS", h2),
        Paragraph("ИТ-стартап (Solana): отчёт по практическим заданиям 1-6", h1),
    ]
    for heading, paragraphs, code in SECTIONS:
        story.append(Paragraph(heading, h2))
        for text in paragraphs:
            story.append(Paragraph(text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), body))
        if code:
            story.append(Preformatted(code, code_style))
            story.append(Spacer(1, 4))
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=42, rightMargin=42, topMargin=36, bottomMargin=36)
    doc.build(story)

if __name__ == "__main__":
    docx_path = ROOT / "Яппаров_Эмир_DS_ИТ_Стартап_Solana_Задания_1-6.docx"
    pdf_path = ROOT / "Яппаров_Эмир_DS_ИТ_Стартап_Solana_Задания_1-6.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print("Built:", docx_path, pdf_path)
